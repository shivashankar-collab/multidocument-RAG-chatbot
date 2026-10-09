import hashlib
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import streamlit as st

try:
    from langchain_core.documents import Document
except ImportError:  # Older LangChain releases
    from langchain.schema import Document

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    from langchain.text_splitter import RecursiveCharacterTextSplitter

try:
    from langchain_huggingface import HuggingFaceEmbeddings, HuggingFaceEndpoint
except ImportError:
    try:
        from langchain_community.embeddings import HuggingFaceEmbeddings
        from langchain_community.llms import HuggingFaceEndpoint
    except ImportError:
        HuggingFaceEmbeddings = None
        HuggingFaceEndpoint = None

try:
    from langchain_community.vectorstores import FAISS
except ImportError:
    FAISS = None

try:
    from langchain_chroma import Chroma
except ImportError:
    try:
        from langchain_community.vectorstores import Chroma
    except ImportError:
        Chroma = None

try:
    from langchain_community.document_loaders import (
        Docx2txtLoader,
        PyPDFLoader,
        UnstructuredPowerPointLoader,
    )
except ImportError:
    Docx2txtLoader = None
    PyPDFLoader = None
    UnstructuredPowerPointLoader = None

try:
    from rank_bm25 import BM25Okapi
except ImportError:
    BM25Okapi = None

try:
    from unstructured.partition.auto import partition
except ImportError:
    partition = None


APP_TITLE = "Multi-Document RAG Chatbot"
DEFAULT_EMBEDDING_MODEL = "sentence-transformers/all-mpnet-base-v2"
DEFAULT_RERANKER_MODEL = "cross-encoder/ms-marco-MiniLM-L-6-v2"
DEFAULT_TOKENIZER_MODEL = "sentence-transformers/all-mpnet-base-v2"
DEFAULT_LLM_MODEL = "Qwen/Qwen3-8B"
DEFAULT_HF_PROVIDER = "nscale"
DEFAULT_CONTEXT_TOKEN_BUDGET = 6_000
READABLE_TEXT_MIN_CHARS = 40
SUPPORTED_FILE_TYPES = [
    "pdf",
    "docx",
    "doc",
    "pptx",
    "ppt",
    "txt",
    "md",
    "csv",
    "tsv",
    "html",
    "htm",
    "xlsx",
    "xls",
    "rtf",
]


@dataclass(frozen=True)
class SearchHit:
    document: Document
    score: float
    vector_score: float
    keyword_score: float
    rerank_score: float | None = None


class KeywordIndex:
    def __init__(self, documents: list[Document]) -> None:
        if BM25Okapi is None:
            raise RuntimeError("rank-bm25 is not installed.")
        self.documents = documents
        self.tokens = [tokenize(doc.page_content) for doc in documents]
        self.index = BM25Okapi(self.tokens)

    def search(self, query: str, limit: int) -> list[tuple[Document, float]]:
        scores = self.index.get_scores(tokenize(query))
        ranked = sorted(enumerate(scores), key=lambda item: item[1], reverse=True)
        return [
            (self.documents[index], float(score))
            for index, score in ranked[:limit]
            if score > 0
        ]


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-zA-Z0-9]+", text.lower())


def content_tokens(text: str) -> set[str]:
    stopwords = {
        "a",
        "an",
        "and",
        "are",
        "as",
        "at",
        "be",
        "by",
        "for",
        "from",
        "has",
        "have",
        "how",
        "i",
        "in",
        "is",
        "it",
        "of",
        "on",
        "or",
        "that",
        "the",
        "this",
        "to",
        "was",
        "what",
        "when",
        "where",
        "which",
        "who",
        "why",
        "with",
    }
    return {token for token in tokenize(text) if len(token) > 2 and token not in stopwords}


@st.cache_resource(show_spinner=False)
def load_hf_tokenizer(model_name: str):
    from transformers import AutoTokenizer

    return AutoTokenizer.from_pretrained(model_name)


def get_tokenizer(model_name: str):
    if not model_name.strip():
        return None
    try:
        return load_hf_tokenizer(model_name)
    except Exception:
        return None


def count_tokens(text: str, tokenizer) -> int:
    if tokenizer is None:
        return max(1, len(text) // 4)
    return len(tokenizer.encode(text, add_special_tokens=False))


def stable_file_id(name: str, content: bytes) -> str:
    digest = hashlib.sha256(content).hexdigest()[:10]
    return f"{Path(name).stem}-{digest}"


def save_upload(uploaded_file) -> Path:
    suffix = Path(uploaded_file.name).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
        handle.write(uploaded_file.getvalue())
        return Path(handle.name)


def clean_text(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def readable_text_length(documents: list[Document]) -> int:
    return sum(len(doc.page_content.strip()) for doc in documents)


def ensure_readable_documents(documents: list[Document], source_name: str) -> list[Document]:
    cleaned = []
    for doc in documents:
        text = clean_text(doc.page_content)
        if text:
            doc.page_content = text
            cleaned.append(doc)

    if readable_text_length(cleaned) < READABLE_TEXT_MIN_CHARS:
        raise RuntimeError(
            f"No readable text was extracted from {source_name}. "
            "If this is a scanned or image-only file, OCR is required."
        )
    return cleaned


def load_plain_text(path: Path, source_name: str, file_id: str) -> list[Document]:
    encodings = ["utf-8", "utf-16", "latin-1"]
    for encoding in encodings:
        try:
            text = path.read_text(encoding=encoding)
            return [
                Document(
                    page_content=text,
                    metadata={
                        "source": source_name,
                        "file_id": file_id,
                        "type": path.suffix.lower().lstrip("."),
                        "extraction_method": "plain_text",
                    },
                )
            ]
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"Could not decode text from {source_name}.")


def load_with_unstructured(path: Path, source_name: str, file_id: str) -> list[Document]:
    if partition is None:
        raise RuntimeError("Generic extraction requires the unstructured package.")

    elements = partition(filename=str(path))
    docs: list[Document] = []
    for index, element in enumerate(elements, start=1):
        text = clean_text(str(element))
        if not text:
            continue

        metadata = {
            "source": source_name,
            "file_id": file_id,
            "type": path.suffix.lower().lstrip("."),
            "element_number": index,
            "extraction_method": "unstructured",
        }
        page_number = getattr(getattr(element, "metadata", None), "page_number", None)
        if page_number:
            metadata["page"] = page_number
        docs.append(Document(page_content=text, metadata=metadata))

    return ensure_readable_documents(docs, source_name)


def load_pdf(path: Path, source_name: str, file_id: str) -> list[Document]:
    if PyPDFLoader is None:
        raise RuntimeError("PyPDFLoader is unavailable. Install langchain-community and pypdf.")
    docs = PyPDFLoader(str(path)).load()
    for doc in docs:
        doc.metadata.update(
            {
                "source": source_name,
                "file_id": file_id,
                "page": doc.metadata.get("page", 0) + 1,
                "type": "pdf",
                "extraction_method": "pypdf",
            }
        )
    return ensure_readable_documents(docs, source_name)


def load_docx(path: Path, source_name: str, file_id: str) -> list[Document]:
    if Docx2txtLoader is None:
        raise RuntimeError("Docx2txtLoader is unavailable. Install langchain-community and docx2txt.")
    docs = Docx2txtLoader(str(path)).load()
    for doc in docs:
        doc.metadata.update(
            {
                "source": source_name,
                "file_id": file_id,
                "type": "docx",
                "extraction_method": "docx2txt",
            }
        )
    return ensure_readable_documents(docs, source_name)


def load_pptx(path: Path, source_name: str, file_id: str) -> list[Document]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for PPTX files.") from exc

    presentation = Presentation(str(path))
    docs: list[Document] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(part.strip() for part in parts if part.strip())
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": source_name,
                        "file_id": file_id,
                        "slide": slide_number,
                        "type": "pptx",
                        "extraction_method": "python-pptx",
                    },
                )
            )
    return ensure_readable_documents(docs, source_name)


def load_ppt(path: Path, source_name: str, file_id: str) -> list[Document]:
    if UnstructuredPowerPointLoader is None:
        raise RuntimeError(
            "Legacy .ppt files require unstructured. Convert to .pptx or install unstructured[ppt]."
        )
    docs = UnstructuredPowerPointLoader(str(path)).load()
    for doc in docs:
        doc.metadata.update(
            {
                "source": source_name,
                "file_id": file_id,
                "type": "ppt",
                "extraction_method": "unstructured_ppt",
            }
        )
    return ensure_readable_documents(docs, source_name)


def load_document(uploaded_file) -> list[Document]:
    source_name = uploaded_file.name
    content = uploaded_file.getvalue()
    file_id = stable_file_id(source_name, content)
    temp_path = save_upload(uploaded_file)
    suffix = temp_path.suffix.lower()

    try:
        loader_errors: list[str] = []
        specialized_loaders = {
            ".pdf": load_pdf,
            ".docx": load_docx,
            ".pptx": load_pptx,
            ".ppt": load_ppt,
        }

        if suffix in {".txt", ".md", ".csv", ".tsv"}:
            return ensure_readable_documents(load_plain_text(temp_path, source_name, file_id), source_name)

        if suffix in specialized_loaders:
            try:
                return specialized_loaders[suffix](temp_path, source_name, file_id)
            except Exception as exc:
                loader_errors.append(str(exc))

        try:
            return load_with_unstructured(temp_path, source_name, file_id)
        except Exception as exc:
            loader_errors.append(str(exc))
            detail = " | ".join(loader_errors)
            raise RuntimeError(
                f"Could not extract readable text from {source_name}. {detail}"
            ) from exc
    finally:
        temp_path.unlink(missing_ok=True)


def split_documents(
    documents: list[Document],
    chunk_size: int,
    chunk_overlap: int,
    tokenizer_model: str,
) -> list[Document]:
    tokenizer = get_tokenizer(tokenizer_model)
    splitter_kwargs = {
        "chunk_size": chunk_size,
        "chunk_overlap": chunk_overlap,
        "separators": ["\n\n", "\n", ". ", " ", ""],
    }
    if tokenizer is not None:
        splitter = RecursiveCharacterTextSplitter.from_huggingface_tokenizer(
            tokenizer,
            **splitter_kwargs,
        )
    else:
        splitter = RecursiveCharacterTextSplitter(**splitter_kwargs)

    chunks = splitter.split_documents(documents)
    for index, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = f"{chunk.metadata.get('file_id', 'doc')}-{index}"
        chunk.metadata["chunk_number"] = index + 1
        chunk.metadata["token_count"] = count_tokens(chunk.page_content, tokenizer)
        chunk.metadata["tokenizer_model"] = tokenizer_model if tokenizer is not None else "estimated"
    return chunks


def make_embeddings(model_name: str):
    if HuggingFaceEmbeddings is None:
        raise RuntimeError("HuggingFaceEmbeddings is unavailable. Install langchain-huggingface.")
    return HuggingFaceEmbeddings(model_name=model_name)


def build_vector_store(chunks: list[Document], embeddings, backend: str):
    if backend == "FAISS":
        if FAISS is None:
            raise RuntimeError("FAISS is unavailable. Install faiss-cpu and langchain-community.")
        return FAISS.from_documents(chunks, embeddings)

    if Chroma is None:
        raise RuntimeError("Chroma is unavailable. Install chromadb and langchain-chroma.")
    return Chroma.from_documents(
        chunks,
        embeddings,
        collection_name=f"rag_{hashlib.sha1(os.urandom(16)).hexdigest()[:8]}",
    )


def normalize_scores(scores: Iterable[float]) -> list[float]:
    values = list(scores)
    if not values:
        return []
    low = min(values)
    high = max(values)
    if high == low:
        return [1.0 for _ in values]
    return [(value - low) / (high - low) for value in values]


def similarity_to_relevance(distance: float) -> float:
    return 1.0 / (1.0 + max(distance, 0.0))


def hybrid_search(
    query: str,
    vector_store,
    keyword_index: KeywordIndex,
    top_k: int,
    vector_weight: float,
) -> list[SearchHit]:
    vector_raw = vector_store.similarity_search_with_score(query, k=max(top_k * 3, 8))
    keyword_raw = keyword_index.search(query, limit=max(top_k * 3, 8))

    by_chunk: dict[str, dict[str, object]] = {}
    vector_scores = [similarity_to_relevance(float(score)) for _, score in vector_raw]
    keyword_scores = [score for _, score in keyword_raw]
    vector_norm = normalize_scores(vector_scores)
    keyword_norm = normalize_scores(keyword_scores)

    for (doc, _), score in zip(vector_raw, vector_norm):
        chunk_id = doc.metadata["chunk_id"]
        by_chunk.setdefault(chunk_id, {"document": doc, "vector": 0.0, "keyword": 0.0})
        by_chunk[chunk_id]["vector"] = max(float(by_chunk[chunk_id]["vector"]), score)

    for (doc, _), score in zip(keyword_raw, keyword_norm):
        chunk_id = doc.metadata["chunk_id"]
        by_chunk.setdefault(chunk_id, {"document": doc, "vector": 0.0, "keyword": 0.0})
        by_chunk[chunk_id]["keyword"] = max(float(by_chunk[chunk_id]["keyword"]), score)

    hits = []
    for entry in by_chunk.values():
        vector_score = float(entry["vector"])
        keyword_score = float(entry["keyword"])
        combined = (vector_weight * vector_score) + ((1.0 - vector_weight) * keyword_score)
        hits.append(
            SearchHit(
                document=entry["document"],
                score=combined,
                vector_score=vector_score,
                keyword_score=keyword_score,
            )
        )
    return sorted(hits, key=lambda hit: hit.score, reverse=True)[:top_k]


def source_label(doc: Document) -> str:
    source = doc.metadata.get("source", "Unknown source")
    if "page" in doc.metadata:
        return f"{source}, page {doc.metadata['page']}"
    if "slide" in doc.metadata:
        return f"{source}, slide {doc.metadata['slide']}"
    return f"{source}, chunk {doc.metadata.get('chunk_number', '?')}"


def format_context(
    hits: list[SearchHit],
    tokenizer_model: str,
    max_context_tokens: int,
) -> str:
    tokenizer = get_tokenizer(tokenizer_model)
    blocks = []
    for number, hit in enumerate(hits, start=1):
        block = f"[{number}] {source_label(hit.document)}\n{hit.document.page_content.strip()}"
        next_context = "\n\n".join([*blocks, block])
        if blocks and count_tokens(next_context, tokenizer) > max_context_tokens:
            break
        blocks.append(block)
    return "\n\n".join(blocks)


@st.cache_resource(show_spinner=False)
def load_reranker(model_name: str):
    from sentence_transformers import CrossEncoder

    return CrossEncoder(model_name)


def rerank_hits(
    question: str,
    hits: list[SearchHit],
    reranker_model: str,
    final_k: int,
) -> tuple[list[SearchHit], str | None]:
    if not reranker_model.strip() or not hits:
        return hits[:final_k], None

    try:
        reranker = load_reranker(reranker_model)
        pairs = [(question, hit.document.page_content) for hit in hits]
        scores = reranker.predict(pairs)
    except Exception as exc:
        return hits[:final_k], str(exc)

    reranked = [
        SearchHit(
            document=hit.document,
            score=hit.score,
            vector_score=hit.vector_score,
            keyword_score=hit.keyword_score,
            rerank_score=float(score),
        )
        for hit, score in zip(hits, scores)
    ]
    return sorted(reranked, key=lambda hit: hit.rerank_score or 0.0, reverse=True)[:final_k], None


def split_sentences(text: str) -> list[str]:
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    return [
        sentence.strip()
        for sentence in re.split(r"(?<=[.!?])\s+|\n+", normalized)
        if len(sentence.strip()) > 20
    ]


def extractive_answer(question: str, hits: list[SearchHit], max_sentences: int = 5) -> str:
    query_tokens = content_tokens(question)
    candidates: list[tuple[float, int, str]] = []

    for citation_number, hit in enumerate(hits, start=1):
        for sentence in split_sentences(hit.document.page_content):
            sentence_tokens = content_tokens(sentence)
            lexical_overlap = len(query_tokens & sentence_tokens)
            semantic_score = hit.rerank_score if hit.rerank_score is not None else hit.score
            score = lexical_overlap + float(semantic_score)
            if lexical_overlap > 0 or citation_number <= 2:
                candidates.append((score, citation_number, sentence))

    if not candidates:
        return "The uploaded documents do not contain enough information to answer this."

    selected: list[tuple[int, str]] = []
    seen: set[str] = set()
    for _, citation_number, sentence in sorted(candidates, key=lambda item: item[0], reverse=True):
        key = sentence.lower()
        if key in seen:
            continue
        seen.add(key)
        selected.append((citation_number, sentence))
        if len(selected) >= max_sentences:
            break

    lines = ["Based on the extracted document text:"]
    for citation_number, sentence in selected:
        lines.append(f"- {sentence} [{citation_number}]")
    return "\n".join(lines)


def fallback_answer(question: str, hits: list[SearchHit], reason: str | None = None) -> str:
    if not hits:
        return "I could not find relevant content in the uploaded documents."

    lines = [extractive_answer(question, hits), ""]
    if reason:
        lines.extend([f"Hosted Hugging Face generation is unavailable, so this answer uses extractive retrieval: {reason}", ""])
    else:
        lines.extend(
            [
                "Set `HUGGINGFACEHUB_API_TOKEN` to enable generative answers. "
                "Without it, answers are extracted directly from the retrieved document text.",
                "",
            ]
        )

    for number, hit in enumerate(hits, start=1):
        snippet = re.sub(r"\s+", " ", hit.document.page_content).strip()
        lines.append(f"[{number}] {snippet[:700]}")
    return "\n\n".join(lines)


def get_huggingface_token() -> str | None:
    token = os.getenv("HUGGINGFACEHUB_API_TOKEN")
    if token:
        return token

    try:
        return st.secrets.get("HUGGINGFACEHUB_API_TOKEN")
    except Exception:
        return None


def generate_answer(
    question: str,
    hits: list[SearchHit],
    model_id: str,
    provider: str,
    tokenizer_model: str,
    max_context_tokens: int,
) -> str:
    token = get_huggingface_token()
    if not token or HuggingFaceEndpoint is None:
        return fallback_answer(question, hits)

    prompt = f"""
You are a document question-answering assistant.
Answer the user's question with respect to the uploaded documents.
Use only the provided context.
Cite every factual claim with bracketed citation numbers like [1] or [2].
Do not use outside knowledge.
Do not guess missing details.
If the context does not contain enough evidence, say: "The uploaded documents do not contain enough information to answer this."
Keep the answer direct and specific.

Context:
{format_context(hits, tokenizer_model, max_context_tokens)}

Question: {question}
Answer:
""".strip()

    try:
        llm = HuggingFaceEndpoint(
            repo_id=model_id,
            provider=provider or None,
            huggingfacehub_api_token=token,
            temperature=0.2,
            max_new_tokens=700,
            task="text-generation",
        )
        response = llm.invoke(prompt)
        return str(response).strip()
    except Exception as exc:
        return fallback_answer(question, hits, str(exc))


@st.cache_resource(show_spinner=False)
def build_indexes(
    file_payloads: tuple[tuple[str, bytes], ...],
    backend: str,
    embedding_model: str,
    tokenizer_model: str,
    chunk_size: int,
    chunk_overlap: int,
):
    documents: list[Document] = []
    errors: list[str] = []

    for name, content in file_payloads:
        uploaded_file = type(
            "UploadedFileProxy",
            (),
            {"name": name, "getvalue": lambda self, data=content: data},
        )()
        try:
            documents.extend(load_document(uploaded_file))
        except Exception as exc:
            errors.append(f"{name}: {exc}")

    if not documents:
        raise RuntimeError("No readable documents were uploaded. " + " ".join(errors))

    chunks = split_documents(documents, chunk_size, chunk_overlap, tokenizer_model)
    embeddings = make_embeddings(embedding_model)
    vector_store = build_vector_store(chunks, embeddings, backend)
    keyword_index = KeywordIndex(chunks)
    return vector_store, keyword_index, chunks, errors


def render_sources(hits: list[SearchHit]) -> None:
    st.subheader("Sources")
    for number, hit in enumerate(hits, start=1):
        with st.expander(
            f"[{number}] {source_label(hit.document)} "
            f"- score {hit.score:.2f}",
            expanded=number <= 2,
        ):
            score_parts = [
                f"Vector {hit.vector_score:.2f}",
                f"Keyword {hit.keyword_score:.2f}",
            ]
            if hit.rerank_score is not None:
                score_parts.append(f"Reranker {hit.rerank_score:.2f}")
            extraction_method = hit.document.metadata.get("extraction_method")
            if extraction_method:
                score_parts.append(f"Extracted by {extraction_method}")
            token_count = hit.document.metadata.get("token_count")
            if token_count:
                score_parts.append(f"{token_count} tokens")
            score_parts.append(f"Chunk {hit.document.metadata.get('chunk_number', '?')}")
            st.caption(" | ".join(score_parts))
            st.write(hit.document.page_content)


def render_extracted_text_preview(chunks: list[Document]) -> None:
    grouped: dict[str, list[Document]] = {}
    for chunk in chunks:
        grouped.setdefault(chunk.metadata.get("source", "Unknown source"), []).append(chunk)

    with st.expander("Extracted text used for answers", expanded=False):
        st.caption(
            "Questions are answered only from this extracted text after retrieval and reranking."
        )
        for source, source_chunks in grouped.items():
            preview = "\n\n".join(chunk.page_content for chunk in source_chunks)
            total_tokens = sum(int(chunk.metadata.get("token_count", 0)) for chunk in source_chunks)
            st.caption(f"{source}: {len(source_chunks)} chunks, about {total_tokens} tokens")
            st.text_area(
                source,
                value=preview[:8_000],
                height=220,
                disabled=True,
            )
            if len(preview) > 8_000:
                st.caption(f"{source}: preview truncated in UI; full extracted text is indexed.")


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="RAG", layout="wide")
    
    # Inject Vercel Speed Insights for performance monitoring
    # This script injects the Speed Insights tracking when deployed on Vercel
    speed_insights_script = """
        <script>
            window.si = window.si || function () { (window.siq = window.siq || []).push(arguments); };
        </script>
        <script defer src="/_vercel/speed-insights/script.js"></script>
    """
    st.components.v1.html(speed_insights_script, height=0)
    
    st.title(APP_TITLE)

    with st.sidebar:
        st.header("Index")
        backend = st.radio("Vector database", ["FAISS", "ChromaDB"], horizontal=True)
        embedding_model = st.text_input("Embedding model", DEFAULT_EMBEDDING_MODEL)
        tokenizer_model = st.text_input("Tokenizer model", DEFAULT_TOKENIZER_MODEL)
        llm_model = st.text_input("Hugging Face LLM", DEFAULT_LLM_MODEL)
        hf_provider = st.text_input("Hugging Face provider", DEFAULT_HF_PROVIDER)
        top_k = st.slider("Sources to retrieve", 2, 12, 5)
        use_reranker = st.checkbox("Use reranker for accuracy", value=True)
        reranker_model = st.text_input("Reranker model", DEFAULT_RERANKER_MODEL)
        vector_weight = st.slider("Vector vs keyword weight", 0.0, 1.0, 0.65, 0.05)
        chunk_size = st.slider("Chunk size in tokens", 200, 1_500, 700, 50)
        chunk_overlap = st.slider("Chunk overlap in tokens", 0, 300, 100, 25)
        max_context_tokens = st.slider(
            "Max context tokens",
            1_000,
            16_000,
            DEFAULT_CONTEXT_TOKEN_BUDGET,
            500,
        )

    uploaded_files = st.file_uploader(
        "Upload documents",
        type=SUPPORTED_FILE_TYPES,
        accept_multiple_files=True,
    )

    if not uploaded_files:
        st.info("Upload one or more documents to build a searchable knowledge base.")
        return

    file_payloads = tuple((file.name, file.getvalue()) for file in uploaded_files)
    selected_backend = "FAISS" if backend == "FAISS" else "Chroma"

    with st.spinner("Reading documents and building hybrid index..."):
        try:
            vector_store, keyword_index, chunks, errors = build_indexes(
                file_payloads,
                selected_backend,
                embedding_model,
                tokenizer_model,
                chunk_size,
                chunk_overlap,
            )
        except Exception as exc:
            st.error(str(exc))
            st.stop()

    st.success(f"Extracted text and indexed {len(uploaded_files)} file(s) into {len(chunks)} chunks.")
    st.caption("Answers are generated only from the extracted document text and cited source chunks.")
    render_extracted_text_preview(chunks)
    for error in errors:
        st.warning(error)

    question = st.chat_input("Ask a question from the uploaded documents")
    if "messages" not in st.session_state:
        st.session_state.messages = []

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if question:
        st.session_state.messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        candidate_k = max(top_k * 4, 16)
        hits = hybrid_search(question, vector_store, keyword_index, candidate_k, vector_weight)
        if use_reranker:
            hits, rerank_error = rerank_hits(question, hits, reranker_model, top_k)
            if rerank_error:
                st.warning(f"Reranker unavailable, using hybrid search order: {rerank_error}")
        else:
            hits = hits[:top_k]

        answer = generate_answer(
            question,
            hits,
            llm_model,
            hf_provider.strip(),
            tokenizer_model,
            max_context_tokens,
        )
        st.session_state.messages.append({"role": "assistant", "content": answer})

        with st.chat_message("assistant"):
            st.markdown(answer)
        render_sources(hits)


if __name__ == "__main__":
    main()
